"""Converts the club's unit library workbook to one CSV of every unit.

Reads docs/complete_library.xlsx (one tab per nation, laid out for printing) and writes
docs/complete_library.csv. docs/library-import.md describes the workbook's layout, the rules
below and the decisions behind them; keep the two in step.

Needs openpyxl. From the repo root:

    python3 -m venv /tmp/wwg-venv && /tmp/wwg-venv/bin/pip install openpyxl
    /tmp/wwg-venv/bin/python scripts/library_csv.py

It prints a check of its totals against the workbook's own Totals tab.
"""

import csv
import re
import sys
from collections import defaultdict
from pathlib import Path

try:
    import openpyxl
    from openpyxl.utils import column_index_from_string
except ImportError:
    sys.exit("openpyxl is missing: see the docstring at the top of this script.")

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "docs" / "complete_library.xlsx"
OUTPUT = ROOT / "docs" / "complete_library.csv"

# Where each nation tab keeps a unit's figures. "label" is the cell beside the count that says
# what it counts (Batt./Sqs./Guns); British and Swedish rows have none.
LAYOUTS = {
    "British": dict(ff="F", count="G", label=None, points="H", status="K", type="L"),
    "French": dict(ff="G", count="H", label="I", points="J", status="M", type="N"),
    "Italy": dict(ff="G", count="H", label="I", points="J", status="M", type="N"),
    "Polish": dict(ff="G", count="H", label="I", points="J", status="M", type="N"),
    "Austrian": dict(ff="F", count="G", label="H", points="I", status="L", type="M"),
    "Prussian": dict(ff="F", count="G", label="H", points="I", status="L", type="M"),
    "Russian": dict(ff="F", count="G", label="H", points="I", status="L", type="M"),
    "Swedish": dict(ff="F", count="G", label=None, points="H", status="K", type="L"),
    "US": dict(ff="F", count="H", label="G", points="I", status="L", type="M"),
}

# The workbook's Totals tab names each army differently from its tab.
TOTALS_NAMES = {
    "British": "British Army",
    "French": "French Army",
    "Italy": "Italian Army",
    "Polish": "Polish Army",
    "Austrian": "Austrian Army",
    "Prussian": "Prussian Army",
    "Russian": "Russian Army",
    "Swedish": "Swedish Army",
    "US": "US Army",
}

STATUSES = {"p": "Painted", "s": "Substitute", "u": "Unpainted"}
CLASSES = {"h": "Heavy", "m": "Medium", "l": "Light", "r": "Light", "s": "Light"}
COUNT_UNITS = {
    "Batt": "Battalions",
    "Sq": "Squadrons",
    "Guns": "Guns",
    "Cos": "Companies",
}
ARM_COUNT_UNITS = {"Infantry": "Battalions", "Cavalry": "Squadrons", "Artillery": "Guns"}

# Rows that start a summary block; units never follow one before the next formation header.
SUMMARY = re.compile(r"Statistics|Totals|^Total |Representing|Men Represented", re.I)
# A trailing "(note)", allowing the stray extra ")" the workbook has in places.
TRAILING_NOTE = re.compile(r"\s*\(([^()]*)\)\)*\s*$")

COLUMNS = [
    "nation",
    "corps",
    "corps_commander",
    "division",
    "division_commander",
    "brigade",
    "brigade_commander",
    "unit",
    "notes",
    "arm",
    "class",
    "type_code",
    "ff",
    "count",
    "count_unit",
    "points",
    "status",
    "status_code",
    "original_name",
    "source_row",
]


def tidy(value):
    return re.sub(r"\s+", " ", value).strip() if isinstance(value, str) else value


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def split_notes(name):
    """Moves trailing bracketed notes out of a name; brackets mid-name are part of it, and so is
    "(bis)", the second of a name."""
    notes, kept = [], ""
    while match := TRAILING_NOTE.search(name):
        if tidy(match.group(1)).lower() == "bis":
            kept = tidy(match.group(0)) + (" " + kept if kept else "")
        else:
            notes.insert(0, tidy(match.group(1)))
        name = name[: match.start()]
    return tidy(f"{name} {kept}"), "; ".join(notes)


def join(*parts):
    return " / ".join(p for p in parts if p)


class Hierarchy:
    """Tracks the corps, division and brigade the rows of a tab are under."""

    def __init__(self):
        self.clear()

    def clear(self):
        self.corps = self.corps_commander = ""
        self.base_corps = self.base_corps_commander = ""
        self.clear_division()
        self.in_summary = True
        self.expect_corps_commander = False
        self.cavalry_corps = False

    def clear_division(self):
        self.division = self.division_commander = ""
        self.clear_brigade()

    def clear_brigade(self):
        self.brigade = self.brigade_commander = ""

    def header(self, text, commander):
        """Applies a non-unit row; returns False when the row isn't one it understands."""
        if text == "Total":
            return True
        if SUMMARY.search(text):
            self.in_summary = True
            self.expect_corps_commander = False
            return True
        if " - Marshall" in text:
            # A French Guard sub-command ("Old Guard Infantry - Marshall Lefebvre"): corps-like.
            name, marshal = (tidy(p) for p in text.split(" - ", 1))
            self.corps = join(self.base_corps, name)
            self.corps_commander = marshal
            self.clear_division()
            self.in_summary = False
            self.expect_corps_commander = False
            return True
        if text in ("Infantry", "Artillery") and self.in_summary is False and not self.brigade:
            return True  # The US tab's totals, which follow its last division unheaded.
        if self.expect_corps_commander:
            if not re.search(r"\b(Division|Divison)\b|Brigade|Corps|Reserve", text) and text != "--":
                line = commander if text == "Commander:" else text
                self.corps_commander = "; ".join(filter(None, [self.corps_commander, line]))
                return True
            self.expect_corps_commander = False
        if text == "Commander:":
            self.corps_commander = commander
            return True
        if re.search(r"Corps (Cavalry|Artillery) Reserve", text):
            # A corps' own reserve (Prussian cavalry, French artillery): division-like.
            self.division, self.division_commander = text.rstrip(":"), commander
            self.clear_brigade()
            self.cavalry_corps = False
            self.in_summary = False
            return True
        if "Cavalry Corps" in text:
            # A Russian cavalry corps inside an infantry corps: division-like.
            self.division, self.division_commander = text, commander
            self.clear_brigade()
            self.cavalry_corps = True
            self.in_summary = False
            return True
        if re.search(r"\bCorps\b|^Cavalry Reserve|^Artillery Reserve", text):
            if text == "Cavalry Reserve (cont.)":
                return True
            self.clear()
            self.corps = self.base_corps = text
            self.corps_commander = commander
            self.in_summary = False
            self.expect_corps_commander = True
            return True
        if re.search(r"\b(Division|Divison)\b", text):
            if self.cavalry_corps and not self.brigade:
                # A division directly under a Russian cavalry corps joins its name.
                self.division = join(self.division, text)
                self.division_commander = join(self.division_commander, commander)
            else:
                self.division, self.division_commander = text, commander
                self.cavalry_corps = False
            self.clear_brigade()
            self.in_summary = False
            return True
        if "Brigade" in text or text == "Independent" or (text == "Cavalry" and not self.in_summary):
            self.brigade, self.brigade_commander = text, commander
            self.in_summary = False
            return True
        if text == "2nd" and self.brigade:
            # A Prussian brigade's second commander.
            self.brigade_commander += f"; 2nd: {commander}"
            return True
        return self.in_summary or not self.corps and not self.division


def subtotal_arm(rows_f, start, layout):
    """The arm of the subtotal row below a unit, for tabs whose rows have no count label."""
    first, last = column_index_from_string(layout["ff"]), column_index_from_string(layout["points"])
    for row in rows_f[start + 1 : start + 30]:
        if is_number(row[first - 1].value) or is_number(row[last - 1].value):
            continue
        for cell in row[first - 1 : last]:
            value = tidy(cell.value)
            if isinstance(value, str) and value.startswith("Batt"):
                return "Infantry"
            if isinstance(value, str) and value.startswith("Sq"):
                return "Cavalry"
    return ""


def arm_of(type_code, label, rows_f, index, layout):
    if not type_code:
        return "Artillery"  # Every row without a type is a battery, troop or gun.
    if label and label.startswith("Batt"):
        return "Infantry"
    if label and label.startswith("Sq"):
        return "Cavalry"
    return subtotal_arm(rows_f, index, layout)  # No label, or "Cos": follow the subtotal.


def read_nation(nation, sheet_f, sheet_v, warnings):
    layout = LAYOUTS[nation]
    col = {k: column_index_from_string(v) for k, v in layout.items() if v}
    width = col["type"] + 1
    rows_f = list(sheet_f.iter_rows(max_col=width))
    rows_v = list(sheet_v.iter_rows(max_col=width))
    hierarchy = Hierarchy()
    units = []
    for index, (row_f, row_v) in enumerate(zip(rows_f, rows_v)):
        texts = [
            tidy(c.value)
            for c in row_f[: col["ff"] - 1]
            if isinstance(c.value, str) and c.value.strip() and not c.value.startswith("=")
        ]
        # A unit row has typed-in figures; subtotal rows hold formulas.
        if not (is_number(row_f[col["ff"] - 1].value) or is_number(row_f[col["points"] - 1].value)):
            if texts and not hierarchy.header(texts[0], texts[1] if len(texts) > 1 else ""):
                warnings.append(f"{nation} row {row_f[0].row}: skipped {texts}")
            continue

        value = lambda key: tidy(row_v[col[key] - 1].value) if key in col else None  # noqa: E731
        original = row_f[[i for i, c in enumerate(row_f) if isinstance(c.value, str)][0]].value.strip()
        name, notes = split_notes(original)
        type_code = value("type") or ""
        status_code = value("status") or ""
        if status_code == "P":
            status_code = "p"  # Decided: a typo for painted.
        label = value("label") if isinstance(value("label"), str) else ""
        arm = arm_of(type_code, label, rows_f, index, layout)
        count_unit = next((u for k, u in COUNT_UNITS.items() if label.startswith(k)), "") if label else ""
        if hierarchy.in_summary:
            warnings.append(f"{nation} row {row_f[0].row}: unit {name!r} outside any formation")
        if not arm:
            warnings.append(f"{nation} row {row_f[0].row}: no arm for {name!r}")
        if status_code not in STATUSES:
            warnings.append(f"{nation} row {row_f[0].row}: status {status_code!r} for {name!r}")
        units.append(
            dict(
                nation=nation,
                corps=hierarchy.corps,
                corps_commander=hierarchy.corps_commander,
                division=hierarchy.division,
                division_commander=hierarchy.division_commander,
                brigade=hierarchy.brigade,
                brigade_commander=hierarchy.brigade_commander,
                unit=name,
                notes=notes,
                arm=arm,
                **{"class": CLASSES.get(type_code, "")},
                type_code=type_code,
                ff=value("ff"),
                count=value("count"),
                count_unit=count_unit or ARM_COUNT_UNITS.get(arm, ""),
                points=value("points"),
                status=STATUSES.get(status_code, ""),
                status_code=status_code,
                original_name=original,
                source_row=row_f[0].row,
            )
        )
    return units


def workbook_totals(sheet):
    """The Totals tab's points and counts per army and arm (its first block of armies)."""
    totals, army = {}, None
    for row in sheet.iter_rows(max_col=3):
        label = tidy(row[0].value)
        if label in ("Infantry", "Cavalry", "Artillery"):
            totals.setdefault((army, label), (row[1].value, row[2].value))
        elif label:
            army = label
    return totals


def main():
    formulas = openpyxl.load_workbook(WORKBOOK)
    values = openpyxl.load_workbook(WORKBOOK, data_only=True)
    warnings, units = [], []
    for nation in LAYOUTS:
        units += read_nation(nation, formulas[nation], values[nation], warnings)

    with OUTPUT.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, COLUMNS)
        writer.writeheader()
        writer.writerows(units)
    print(f"Wrote {len(units)} units to {OUTPUT.relative_to(ROOT)}")

    print("\nAgainst the workbook's Totals tab (points, count):")
    ours = defaultdict(lambda: [0, 0])
    for u in units:
        ours[(u["nation"], u["arm"])][0] += u["points"]
        ours[(u["nation"], u["arm"])][1] += u["count"]
    expected = workbook_totals(values["Totals"])
    for nation in LAYOUTS:
        for arm in ("Infantry", "Cavalry", "Artillery"):
            theirs = expected.get((TOTALS_NAMES[nation], arm))
            mine = tuple(ours[(nation, arm)])
            flag = "ok" if theirs == mine else f"workbook {theirs}"
            print(f"  {nation:9} {arm:9} {mine}  {flag}")

    if warnings:
        print("\nWarnings:")
        for w in warnings:
            print(f"  {w}")


if __name__ == "__main__":
    main()
